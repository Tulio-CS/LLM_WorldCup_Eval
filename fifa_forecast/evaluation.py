"""Compare actual results against the AI forecasts (read-only).

Joins ``forecast_runs`` (successful, valid-JSON predictions) to ``match_results``
(``status='finished'``) by ``match_id`` and scores each prediction:

* **outcome_correct** — did the model pick the right winner / draw?
* **exact_score_correct** — did it nail the exact scoreline?
* **abs_goal_diff_error** / **total_goals_error** — numeric closeness.
* **brier** — 3-way probability Brier score (probability/six-hats prompts).

Team-order handling: predictions are stored in the *prompt* presentation order
(``prompt_team_1`` vs ``prompt_team_2``); a ``reversed`` run has them swapped
relative to the canonical ``team_1``/``team_2``. Everything here is mapped back
to canonical order before being compared to the actual (canonical) result, so
original and reversed runs are directly comparable.
"""

from __future__ import annotations

import sqlite3
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

from . import config as cfg
from .prompts import PROMPTS_WITH_PROBABILITIES


@dataclass
class Evaluation:
    tables: dict[str, pd.DataFrame] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)


def load_joined(db_path: str | Path) -> pd.DataFrame:
    path = Path(db_path)
    if not path.exists():
        raise FileNotFoundError(f"Database not found: {path}")
    con = sqlite3.connect(path)
    try:
        preds = pd.read_sql_query(
            "SELECT * FROM forecast_runs "
            "WHERE execution_status='success' AND json_valid=1",
            con,
        )
        try:
            res = pd.read_sql_query(
                "SELECT match_id, status, actual_score_team_1, actual_score_team_2, "
                "actual_winner FROM match_results WHERE status='finished' "
                "AND actual_score_team_1 IS NOT NULL AND actual_score_team_2 IS NOT NULL",
                con,
            )
        except Exception:
            res = pd.DataFrame()
    finally:
        con.close()
    if preds.empty or res.empty:
        return pd.DataFrame()
    return preds.merge(res, on="match_id", how="inner", suffixes=("", "_res"))


def _canonicalize(df: pd.DataFrame) -> pd.DataFrame:
    """Add canonical-order predicted scores/probabilities + correctness columns."""
    df = df.copy()
    reversed_mask = df["team_order_type"] == "reversed"

    s1 = pd.to_numeric(df["parsed_score_team_1"], errors="coerce")
    s2 = pd.to_numeric(df["parsed_score_team_2"], errors="coerce")
    # canonical: for reversed runs, prompt_team_1 is the canonical team_2
    df["pred_c1"] = np.where(reversed_mask, s2, s1)
    df["pred_c2"] = np.where(reversed_mask, s1, s2)

    p1 = pd.to_numeric(df["parsed_team1_win_probability"], errors="coerce")
    pd_ = pd.to_numeric(df["parsed_draw_probability"], errors="coerce")
    p2 = pd.to_numeric(df["parsed_team2_win_probability"], errors="coerce")
    df["pc_team1"] = np.where(reversed_mask, p2, p1)
    df["pc_draw"] = pd_
    df["pc_team2"] = np.where(reversed_mask, p1, p2)

    a1 = pd.to_numeric(df["actual_score_team_1"], errors="coerce")
    a2 = pd.to_numeric(df["actual_score_team_2"], errors="coerce")
    df["actual_c1"] = a1
    df["actual_c2"] = a2

    def outcome(c1: pd.Series, c2: pd.Series) -> pd.Series:
        return np.where(c1 > c2, "team_1", np.where(c1 < c2, "team_2", "draw"))

    df["pred_outcome"] = outcome(df["pred_c1"], df["pred_c2"])
    df["actual_outcome"] = df["actual_winner"].fillna(
        pd.Series(outcome(a1, a2), index=df.index)
    )

    df["outcome_correct"] = (df["pred_outcome"] == df["actual_outcome"]).astype(int)
    df["exact_score_correct"] = (
        (df["pred_c1"] == a1) & (df["pred_c2"] == a2)
    ).astype(int)
    df["abs_goal_diff_error"] = ((df["pred_c1"] - df["pred_c2"]) - (a1 - a2)).abs()
    df["total_goals_error"] = ((df["pred_c1"] + df["pred_c2"]) - (a1 + a2)).abs()

    # 3-way Brier (probability prompts only); normalize probs to sum 1.
    has_prob = df["prompt_id"].isin(PROMPTS_WITH_PROBABILITIES)
    psum = (df["pc_team1"] + df["pc_draw"] + df["pc_team2"]).replace(0, np.nan)
    pt1 = df["pc_team1"] / psum
    ptd = df["pc_draw"] / psum
    pt2 = df["pc_team2"] / psum
    y1 = (df["actual_outcome"] == "team_1").astype(float)
    yd = (df["actual_outcome"] == "draw").astype(float)
    y2 = (df["actual_outcome"] == "team_2").astype(float)
    brier = (pt1 - y1) ** 2 + (ptd - yd) ** 2 + (pt2 - y2) ** 2
    df["brier"] = np.where(has_prob & psum.notna(), brier, np.nan)
    return df


def _agg(df: pd.DataFrame, by: list[str]) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()
    g = df.groupby(by, dropna=False)
    out = g.agg(
        n=("run_id", "size"),
        outcome_acc=("outcome_correct", "mean"),
        exact_acc=("exact_score_correct", "mean"),
        mean_abs_gd_err=("abs_goal_diff_error", "mean"),
        mean_total_goals_err=("total_goals_error", "mean"),
        mean_brier=("brier", "mean"),
    ).reset_index()
    for col in ("outcome_acc", "exact_acc"):
        out[col] = (out[col] * 100).round(1)  # as percentages
    for col in ("mean_abs_gd_err", "mean_total_goals_err"):
        out[col] = out[col].round(3)
    out["mean_brier"] = out["mean_brier"].round(4)
    return out


def build_evaluation(config: cfg.Config) -> Evaluation:
    ev = Evaluation()
    joined = load_joined(cfg.ROOT / config.database_path)
    if joined.empty:
        ev.notes.append(
            "No finished results joined to valid predictions yet. Ingest results "
            "(fetch / add-results) for matches that already happened, then re-run."
        )
        return ev

    # Tolerate older databases that predate some columns.
    for col, default in (("match_moment", None), ("team_order_type", "original")):
        if col not in joined.columns:
            joined[col] = default

    df = _canonicalize(joined)

    leaderboard = _agg(df, ["provider", "model"]).sort_values(
        ["outcome_acc", "mean_brier"], ascending=[False, True]
    ).reset_index(drop=True)
    ev.tables["Leaderboard (by model)"] = leaderboard
    ev.tables["By Model & Prompt"] = _agg(df, ["model", "prompt_id"])
    if df["match_moment"].notna().any() and df["match_moment"].nunique(dropna=True) > 1:
        ev.tables["By Moment"] = _agg(df, ["match_moment", "model"])
    ev.tables["By Match"] = _agg(df, ["match_id", "team_1", "team_2"])

    # Per-prediction detail (canonical), trimmed to the useful columns.
    detail_cols = [
        "run_id", "match_id", "team_1", "team_2", "match_moment", "team_order_type",
        "provider", "model", "prompt_id", "repetition_number",
        "pred_c1", "pred_c2", "actual_c1", "actual_c2",
        "pred_outcome", "actual_outcome", "outcome_correct", "exact_score_correct",
        "abs_goal_diff_error", "total_goals_error", "brier",
    ]
    ev.tables["Per Prediction"] = df[[c for c in detail_cols if c in df.columns]].copy()

    ev.notes.append(
        f"{df['match_id'].nunique()} finished match(es) · {len(df)} predictions evaluated."
    )
    ev.notes.append(
        "Scores/probabilities are mapped to canonical team order before scoring, "
        "so original and reversed runs are comparable. outcome_acc/exact_acc are %, "
        "mean_brier is the 3-way Brier (lower is better; probability prompts only)."
    )
    return ev


def write_evaluation_excel(ev: Evaluation, output_path: str | Path | None = None) -> Path:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    target = (
        Path(output_path)
        if output_path
        else (cfg.EXPORTS_DIR / "FIFA_Evaluation_Report.xlsx")
    )
    target.parent.mkdir(parents=True, exist_ok=True)

    def clean(df: pd.DataFrame) -> pd.DataFrame:
        if df.empty:
            return df
        return df.map(
            lambda v: ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v
        )

    notes_df = pd.DataFrame({"notes": ev.notes or ["(no data)"]})
    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        clean(notes_df).to_excel(writer, sheet_name="Notes", index=False)
        for name, df in ev.tables.items():
            clean(df).to_excel(writer, sheet_name=name[:31], index=False)
    return target


def _fmt(df: pd.DataFrame) -> str:
    if df is None or df.empty:
        return "  (no data)"
    with pd.option_context("display.width", 200, "display.max_columns", None):
        return df.to_string(index=False)


def print_evaluation(ev: Evaluation) -> None:
    print("=" * 74)
    print("AI FORECAST vs ACTUAL RESULTS")
    print("=" * 74)
    if not ev.tables:
        for n in ev.notes:
            print(f"- {n}")
        return
    print("\nLeaderboard (by model):")
    print(_fmt(ev.tables.get("Leaderboard (by model)")))
    print("\nBy model & prompt:")
    print(_fmt(ev.tables.get("By Model & Prompt")))
    if "By Moment" in ev.tables:
        print("\nBy moment:")
        print(_fmt(ev.tables.get("By Moment")))
    print("\nNotes:")
    for n in ev.notes:
        print(f"- {n}")
